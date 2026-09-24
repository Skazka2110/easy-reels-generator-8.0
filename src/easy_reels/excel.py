from __future__ import annotations

import os
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .errors import WorkbookError
from .models import HookEditorRow, HookRow


CANONICAL_HEADERS = ("Хук", "Подхук", "Описание", "Готовый файл", "Статус")
HEADER_KEYS = {header.casefold(): header for header in CANONICAL_HEADERS}


def _plain_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


class HooksWorkbook:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        try:
            self.book = load_workbook(self.path, data_only=False)
        except FileNotFoundError as exc:
            raise WorkbookError(f"Файл hooks.xlsx не найден: {self.path}") from exc
        except PermissionError as exc:
            raise WorkbookError(
                "Не удалось открыть список хуков. Закройте hooks.xlsx в другой программе и повторите попытку."
            ) from exc
        except Exception as exc:
            raise WorkbookError(f"Не удалось прочитать hooks.xlsx: {exc}") from exc

        self.dirty = False
        self.sheet = self._select_sheet()
        self.columns = self._ensure_columns()

    def _headers_for(self, sheet: Worksheet) -> dict[str, int]:
        result: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(1, column).value
            if value is None:
                continue
            normalized = str(value).strip().casefold()
            if normalized in HEADER_KEYS:
                if normalized in result:
                    raise WorkbookError(
                        f"На листе «{sheet.title}» заголовок «{HEADER_KEYS[normalized]}» указан несколько раз."
                    )
                result[normalized] = column
        return result

    def _select_sheet(self) -> Worksheet:
        candidates: list[Worksheet] = []
        for sheet in self.book.worksheets:
            if sheet.sheet_state != "visible":
                continue
            if "хук" in self._headers_for(sheet):
                candidates.append(sheet)
        if not candidates:
            raise WorkbookError(
                "Не найден видимый лист с заголовком «Хук» в первой строке."
            )
        if len(candidates) > 1:
            names = ", ".join(f"«{sheet.title}»" for sheet in candidates)
            raise WorkbookError(
                f"Найдено несколько подходящих листов ({names}). Оставьте заголовок «Хук» только на одном рабочем листе."
            )
        return candidates[0]

    def _ensure_columns(self) -> dict[str, int]:
        columns = self._headers_for(self.sheet)
        for canonical in CANONICAL_HEADERS:
            key = canonical.casefold()
            if key not in columns:
                column = self.sheet.max_column + 1
                self.sheet.cell(1, column, canonical)
                columns[key] = column
                self.dirty = True
        return columns

    def _cell(self, row: int, header: str):
        return self.sheet.cell(row, self.columns[header.casefold()])

    def _set_status(self, row: int, message: str) -> None:
        self._cell(row, "Статус").value = message[:500]
        self.dirty = True

    def pending_rows(self, *, record_errors: bool = True) -> list[HookRow]:
        pending: list[HookRow] = []
        for row in range(2, self.sheet.max_row + 1):
            ready = _plain_text(self._cell(row, "Готовый файл").value)
            if ready:
                continue

            hook_cell = self._cell(row, "Хук")
            subhook_cell = self._cell(row, "Подхук")
            hook = _plain_text(hook_cell.value)
            subhook = _plain_text(subhook_cell.value)

            if not hook:
                if subhook and record_errors:
                    self._set_status(row, "Ошибка: подхук заполнен без хука")
                continue

            if hook_cell.data_type == "f" or subhook_cell.data_type == "f":
                if record_errors:
                    self._set_status(
                        row, "Ошибка: формулы в хуке и подхуке не поддерживаются"
                    )
                continue

            pending.append(
                HookRow(
                    row_number=row,
                    hook=hook,
                    subhook=subhook,
                    description=_plain_text(self._cell(row, "Описание").value),
                )
            )
        return pending

    def next_pending(self, *, record_errors: bool = True) -> HookRow | None:
        rows = self.pending_rows(record_errors=record_errors)
        return rows[0] if rows else None

    def editor_rows(self) -> list[HookEditorRow]:
        """Return all meaningful rows for the built-in editor."""
        rows: list[HookEditorRow] = []
        for row in range(2, self.sheet.max_row + 1):
            values = HookEditorRow(
                hook=_plain_text(self._cell(row, "Хук").value),
                subhook=_plain_text(self._cell(row, "Подхук").value),
                description=_plain_text(self._cell(row, "Описание").value),
                ready_file=_plain_text(self._cell(row, "Готовый файл").value),
                status=_plain_text(self._cell(row, "Статус").value),
            )
            if any(
                (
                    values.hook,
                    values.subhook,
                    values.description,
                    values.ready_file,
                    values.status,
                )
            ):
                rows.append(values)
        return rows

    def replace_editor_rows(self, rows: list[HookEditorRow]) -> None:
        """Replace worksheet contents while preserving the canonical header row."""
        if self.sheet.max_row > 1:
            self.sheet.delete_rows(2, self.sheet.max_row - 1)
        for row_number, item in enumerate(rows, start=2):
            self._cell(row_number, "Хук").value = item.hook
            self._cell(row_number, "Подхук").value = item.subhook
            self._cell(row_number, "Описание").value = item.description
            self._cell(row_number, "Готовый файл").value = item.ready_file
            self._cell(row_number, "Статус").value = item.status
        self.columns = self._ensure_columns()
        self.dirty = True

    def mark_success(self, row: int, file_name: str) -> None:
        self._cell(row, "Готовый файл").value = file_name
        self._set_status(row, "Готово")

    def mark_error(self, row: int, message: str) -> None:
        one_line = " ".join(str(message).split())
        self._set_status(row, f"Ошибка: {one_line}")

    def create_backup(self, backup_dir: str | Path) -> Path:
        backup_dir = Path(backup_dir)
        backup_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        target = backup_dir / f"hooks_{timestamp}.xlsx"
        counter = 2
        while target.exists():
            target = backup_dir / f"hooks_{timestamp}_{counter}.xlsx"
            counter += 1
        try:
            shutil.copy2(self.path, target)
        except OSError as exc:
            raise WorkbookError(f"Не удалось создать резервную копию hooks.xlsx: {exc}") from exc
        return target

    def save(self) -> None:
        temp = self.path.with_name(
            f".{self.path.stem}.erg-{uuid.uuid4().hex}{self.path.suffix}"
        )
        try:
            self.book.save(temp)
            os.replace(temp, self.path)
        except PermissionError as exc:
            raise WorkbookError(
                "Не удалось сохранить список хуков. Закройте hooks.xlsx в другой программе и повторите попытку."
            ) from exc
        except OSError as exc:
            raise WorkbookError(f"Не удалось безопасно сохранить hooks.xlsx: {exc}") from exc
        finally:
            try:
                if temp.exists():
                    temp.unlink()
            except OSError:
                pass
        self.dirty = False


def create_workbook_template(path: str | Path) -> Path:
    path = Path(path)
    if path.exists():
        raise WorkbookError(f"Файл уже существует: {path}")
    book = Workbook()
    sheet = book.active
    sheet.title = "Хуки"
    sheet.append(CANONICAL_HEADERS)
    sheet.freeze_panes = "A2"
    widths = {"A": 52, "B": 44, "C": 34, "D": 42, "E": 34}
    fill = PatternFill("solid", fgColor="C6EFCE")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="006D3C")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path
